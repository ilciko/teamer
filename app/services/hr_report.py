"""Report per HR delle proposte di aumento: export Excel (contenuto = pagina stampabile)."""
import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Font

from app import db

HEADERS = ["Dipendente", "Team", "Livello attuale", "Nuovo livello",
           "RAL attuale €", "Aumento €", "Aumento %", "Nuova RAL €",
           "Decorrenza", "Stato", "Nota"]


def report_rows(budget_id, statuses: tuple[str, ...]) -> list[dict]:
    where, params = "pr.voided=0", []
    where += " AND pr.status IN (%s)" % ",".join("?" * len(statuses))
    params.extend(statuses)
    if budget_id:
        where += " AND pr.budget_id=?"
        params.append(budget_id)
    rows = db.query(
        f"""SELECT pr.*, p.full_name, t.name AS team_name,
                   (SELECT amount FROM ral_event r WHERE r.person_id=pr.person_id AND r.voided=0
                    ORDER BY r.effective_date DESC, r.id DESC LIMIT 1) AS ral,
                   (SELECT livello FROM employment e WHERE e.person_id=pr.person_id
                    AND e.voided=0 AND e.end_date IS NULL ORDER BY e.start_date DESC LIMIT 1) AS livello
            FROM ral_proposal pr
            JOIN person p ON p.id = pr.person_id
            LEFT JOIN person_assignment pa ON pa.person_id=p.id AND pa.voided=0 AND pa.end_date IS NULL
            LEFT JOIN team t ON t.id = pa.team_id
            WHERE {where} ORDER BY p.full_name, pr.effective_date""", tuple(params))
    out = []
    for r in rows:
        ral = r["ral"] or 0
        out.append({**dict(r),
                    "pct": (r["amount"] / ral * 100) if ral and r["amount"] else 0,
                    "new_ral": round(ral + (r["amount"] or 0), 2)})
    return out


def build_xlsx(budget, rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Proposte aumenti"
    title = f"Proposte di aumento — {budget['period']}" if budget else "Proposte di aumento"
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"Generato il {date.today().strftime('%d/%m/%Y')}"])
    ws.append([])
    ws.append(HEADERS)
    for c in ws[4]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([
            r["full_name"], r["team_name"] or "",
            db.lookup_label("livello", r["livello"]) if r["livello"] else "",
            db.lookup_label("livello", r["new_livello"]) if r["new_livello"] else "",
            r["ral"] or 0, r["amount"] or 0, round(r["pct"], 2), r["new_ral"],
            r["effective_date"] or "", db.lookup_label("proposal_status", r["status"]),
            r["note"] or "",
        ])
    total_row = ["Totale", "", "", "", "", sum(r["amount"] or 0 for r in rows), "", "", "", "", ""]
    ws.append(total_row)
    ws[ws.max_row][0].font = Font(bold=True)
    ws[ws.max_row][5].font = Font(bold=True)
    if budget:
        ws.append([])
        ws.append(["Budget", budget["period"], "tetto", budget["total_cap"]])
    widths = [24, 16, 13, 13, 13, 12, 11, 13, 12, 13, 28]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=4, column=i).column_letter].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
