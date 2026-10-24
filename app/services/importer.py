"""Import persone da Excel (HR, 0.44): censimento dello stato attuale, con tutti i campi della persona.

Ogni riga nasce intestata al **director** indicato (`person.owner_user_id`); il team è facoltativo e,
se c'è, è un team di quel director (creato se manca). Contratto, RAL e codice di listino sono i campi
che HR registra dalla scheda; qui arrivano in blocco. Il riporto diretto non si importa: si calcola da
team e ruolo (lo fa la route dopo l'import).
"""
import io
from datetime import date, datetime

from openpyxl import Workbook, load_workbook

from app import db

COLUMNS = [
    ("nome", "Nome e cognome (obbligatorio)"),
    ("director", "Director titolare: nome utente o nome visualizzato (obbligatorio)"),
    ("email", "Email"),
    ("data_assunzione", "Data assunzione (GG/MM/AAAA)"),
    ("data_nascita", "Data nascita (GG/MM/AAAA)"),
    ("team", "Team del director (facoltativo; creato se non esiste)"),
    ("ruolo", "Ruolo: ic / tl / mgr / sr (con il team)"),
    ("codice_listino", "Codice di listino (Azienda → Listino)"),
    ("livello", "Livello (codice della lista aziendale; uno nuovo viene aggiunto)"),
    ("ccnl", "CCNL (vuoto = predefinito aziendale)"),
    ("contratto", "Contratto: indeterminato / determinato / apprendistato / collaborazione"),
    ("scadenza_contratto", "Scadenza contratto se a termine (GG/MM/AAAA)"),
    ("smartwork", "Smart working: no / parziale / full"),
    ("timbratura", "Timbratura: esente / 4_timbrate / solo_presenza"),
    ("part_time_pct", "Part-time % (100 = full time)"),
    ("ore_settimanali", "Ore settimanali"),
    ("ral", "RAL attuale €"),
    ("note", "Note"),
]


def build_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Persone"
    for col, (key, _) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=key)
        cell.font = cell.font.copy(bold=True)
        ws.column_dimensions[cell.column_letter].width = max(14, len(key) + 4)
    info = wb.create_sheet("Istruzioni")
    info.column_dimensions["A"].width = 30
    info.column_dimensions["B"].width = 70
    info.append(["Colonna", "Descrizione"])
    for key, desc in COLUMNS:
        info.append([key, desc])
    info.append([])
    info.append(["Note", "Una riga per persona. Obbligatori 'nome' e 'director'; le altre colonne possono restare vuote."])
    info.append(["", "Le date valgono sia come data Excel sia come testo GG/MM/AAAA o AAAA-MM-GG."])
    info.append(["", "L'import è un censimento: crea lo stato attuale (assegnazione se c'è il team, contratto, RAL con tipo 'censimento')."])
    info.append(["", "Senza team la persona è del director ma non è ancora in organigramma: il director la assegna dalla scheda."])
    info.append(["", "Persone già presenti con lo stesso nome vengono saltate."])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _parse_date(v) -> str | None:
    if v in (None, ""):
        return None
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    s = str(v).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    raise ValueError(f"data non riconosciuta: {s!r}")


def _norm(v) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def _num(v) -> float | None:
    if v in (None, ""):
        return None
    return float(str(v).strip().replace(",", "."))


def _ensure_lookup(category: str, code: str) -> None:
    """Un livello nuovo dal foglio entra nella lista aziendale (scritta dalla connessione di sistema)."""
    if not db.one("SELECT 1 FROM lookup WHERE category=? AND code=?", (category, code)):
        db.sys.execute(
            "INSERT INTO lookup (category, code, label, sort_order) VALUES (?,?,?,99)",
            (category, code, code))


def directors() -> dict[str, int]:
    """{nome utente o nome visualizzato (minuscolo): id} dei director attivi."""
    out = {}
    for u in db.sys.query("SELECT id, username, display_name FROM app_user WHERE role='director' AND active=1"):
        out[u["username"].lower()] = u["id"]
        if u["display_name"]:
            out[u["display_name"].strip().lower()] = u["id"]
    return out


def import_workbook(content: bytes, ref_date: str) -> dict:
    """Importa il foglio 'Persone'. Ritorna {created, skipped, errors: [(riga, msg)], with_team: [id]}."""
    wb = load_workbook(io.BytesIO(content), data_only=True)
    ws = wb["Persone"] if "Persone" in wb.sheetnames else wb.active
    headers = [(_norm(c.value) or "").lower() for c in ws[1]]
    dirs = directors()

    def col(row, key):
        try:
            return row[headers.index(key)]
        except ValueError:
            return None

    created, skipped, errors, with_team = [], [], [], []
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        name = _norm(col(row, "nome"))
        if not name:
            continue
        try:
            if db.one("SELECT id FROM person WHERE lower(full_name)=lower(?)", (name,)):
                skipped.append(name)
                continue
            d = (_norm(col(row, "director")) or "").lower()
            if d not in dirs:
                raise ValueError(f"director {d!r} inesistente o non attivo" if d else "director mancante")
            owner = dirs[d]
            hire = _parse_date(col(row, "data_assunzione"))
            birth = _parse_date(col(row, "data_nascita"))
            start = hire or ref_date
            team_name = _norm(col(row, "team"))
            lp = None
            code = _norm(col(row, "codice_listino"))
            if team_name and code:
                row_lp = db.one("SELECT id FROM listino_profile WHERE upper(code)=upper(?)", (code,))
                if not row_lp:
                    raise ValueError(f"codice di listino {code!r} inesistente (Azienda → Listino)")
                lp = row_lp["id"]
            livello = _norm(col(row, "livello"))
            if livello and livello.endswith(".0"):
                livello = livello[:-2]
            contratto = (_norm(col(row, "contratto")) or "").lower() or None
            pt = _num(col(row, "part_time_pct"))
            ore = _num(col(row, "ore_settimanali"))
            scadenza = _parse_date(col(row, "scadenza_contratto"))
            ral = _num(col(row, "ral"))
            team_id = None
            if team_name:
                # un team del director (mai archiviato); se manca si crea, suo
                t = db.one("SELECT id FROM team WHERE lower(name)=lower(?) AND archived_at IS NULL AND owner_user_id=?",
                           (team_name, owner))
                team_id = t["id"] if t else db.sys.execute(
                    "INSERT INTO team (name, owner_user_id) VALUES (?,?)", (team_name, owner))
            if livello:
                _ensure_lookup("livello", livello)
            # la riga è tutta o niente: una riga sbagliata non lascia una persona a metà
            with db.transaction() as conn:
                pid = conn.execute(
                    "INSERT INTO person (full_name, email, hire_date, birth_date, notes, owner_user_id) VALUES (?,?,?,?,?,?)",
                    (name, _norm(col(row, "email")), hire, birth, _norm(col(row, "note")), owner)).lastrowid
                if team_id:
                    conn.execute(
                        "INSERT INTO person_assignment (person_id, team_id, role_type, start_date, listino_profile_id) "
                        "VALUES (?,?,?,?,?)", (pid, team_id, (_norm(col(row, "ruolo")) or "ic").lower(), start, lp))
                if any([livello, contratto, _norm(col(row, "ccnl")), _norm(col(row, "smartwork")),
                        _norm(col(row, "timbratura")), pt, ore]):
                    conn.execute(
                        "INSERT INTO employment (person_id, livello, ccnl, tipo_contratto, smartwork, "
                        "timbratura_scheme, part_time_pct, weekly_hours, start_date, end_date) VALUES (?,?,?,?,?,?,?,?,?,?)",
                        (pid, livello, _norm(col(row, "ccnl")) or db.get_meta("default_ccnl") or None, contratto,
                         (_norm(col(row, "smartwork")) or "").lower() or None,
                         (_norm(col(row, "timbratura")) or "").lower() or None,
                         pt if pt is not None else 100.0, ore, start, scadenza))
                if ral is not None:
                    # chi è assunto dopo il censimento parte dalla sua assunzione, non dalla data dell'import
                    conn.execute(
                        "INSERT INTO ral_event (person_id, effective_date, amount, delta, change_type, reason) "
                        "VALUES (?,?,?,0,'censimento','import Excel')",
                        (pid, max(hire, ref_date) if hire else ref_date, ral))
            if team_id:
                with_team.append(pid)
            created.append(name)
        except Exception as exc:  # riga malformata: si salta, le altre proseguono
            errors.append((i, f"{name}: {exc}"))
    return {"created": created, "skipped": skipped, "errors": errors, "with_team": with_team}
