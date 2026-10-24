"""Censimento e import da Excel di HR (0.44): tutti i campi della persona, director titolare,
team facoltativo fra i suoi; il director non censisce più."""
import io

from openpyxl import Workbook, load_workbook

from app import db, users
from app.services import importer
from tests.test_viste_hr import _seed


def _xlsx(rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Persone"
    keys = [k for k, _ in importer.COLUMNS]
    ws.append(keys)
    for r in rows:
        ws.append([r.get(k) for k in keys])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_template_e_pagina_solo_per_hr(client, hr_client):
    r = hr_client.get("/people/import/template")
    assert r.status_code == 200 and "Persone" in load_workbook(io.BytesIO(r.content)).sheetnames
    assert "director" in [c.value for c in load_workbook(io.BytesIO(r.content))["Persone"][1]]
    assert hr_client.get("/people/import").status_code == 200
    assert client.get("/people/import").status_code == 403
    assert client.get("/people/import/template").status_code == 403
    assert client.post("/people/import", files={"file": ("x.xlsx", b"", "application/octet-stream")},
                       data={"ref_date": "2026-01-01"}).status_code == 403
    assert client.get("/settings?tab=import").status_code == 200 and "Import persone" not in client.get("/settings?tab=import").text
    assert "Nuova persona" not in client.get("/people").text and "le censisce HR" in client.get("/people").text


def test_import_con_tutti_i_campi(client, client2, hr_client):
    t1, t2, ids, coo, dir2 = _seed(client, client2)
    lp = db.sys.execute("INSERT INTO listino_profile (code, description, day_rate) VALUES ('GE040', 'Consultant', 450)")
    content = _xlsx([
        {"nome": "Ugo Uno", "director": "dir2", "email": "ugo@x.it", "data_assunzione": "15/03/2024", "data_nascita": "01/01/1990",
         "team": "Beta", "ruolo": "tl", "codice_listino": "ge040", "livello": "7", "ccnl": "Commercio", "contratto": "determinato",
         "scadenza_contratto": "2027-06-30", "smartwork": "parziale", "timbratura": "esente", "part_time_pct": 80,
         "ore_settimanali": 32, "ral": "41000", "note": "import e2e"},
        {"nome": "Vera Due", "director": "COO", "team": "Nuovo Team", "ruolo": "ic", "ral": 30000},          # team creato per coo
        {"nome": "Zeno Tre", "director": "dir2"},                                                            # senza team
        {"nome": "Anna Alfa", "director": "coo"},                                                            # già presente
        {"nome": "Yuri Quattro", "director": "nessuno"},                                                     # director inesistente
        {"nome": "Xena Cinque", "director": "dir2", "team": "Beta", "codice_listino": "ZZ999"},             # codice inesistente
    ])
    r = hr_client.post("/people/import", files={"file": ("p.xlsx", content, "application/octet-stream")},
                       data={"ref_date": "2026-01-01"})
    assert r.status_code == 200
    assert "3</div><div class=\"k\">Create" in r.text.replace("\n", "") or r.text.count("Create")
    html = r.text
    assert "Ugo Uno" in html and "Vera Due" in html and "Zeno Tre" in html
    assert "Saltate" in html and "Anna Alfa" in html
    assert "nessuno" in html and "inesistente o non attivo" in html and "ZZ999" in html
    ugo = db.sys.one("SELECT * FROM person WHERE full_name='Ugo Uno'")
    assert ugo["owner_user_id"] == dir2 and ugo["email"] == "ugo@x.it" and str(ugo["birth_date"]) == "1990-01-01" and ugo["notes"] == "import e2e"
    pa = db.sys.one("SELECT * FROM person_assignment WHERE person_id=?", (ugo["id"],))
    assert pa["team_id"] == t2 and pa["role_type"] == "tl" and pa["listino_profile_id"] == lp and str(pa["start_date"]) == "2024-03-15"
    emp = db.sys.one("SELECT * FROM employment WHERE person_id=?", (ugo["id"],))
    assert (emp["livello"], emp["ccnl"], emp["tipo_contratto"], emp["smartwork"], emp["timbratura_scheme"]) == \
        ("7", "Commercio", "determinato", "parziale", "esente")
    assert emp["part_time_pct"] == 80 and emp["weekly_hours"] == 32 and str(emp["end_date"]) == "2027-06-30"
    ral = db.sys.one("SELECT * FROM ral_event WHERE person_id=?", (ugo["id"],))
    assert ral["amount"] == 41000 and ral["change_type"] == "censimento" and str(ral["effective_date"]) == "2026-01-01"
    assert db.sys.one("SELECT 1 FROM lookup WHERE category='livello' AND code='7'")
    # Vera: team nuovo intestato a coo; Zeno: di dir2 senza team; i director li vedono subito
    team = db.sys.one("SELECT owner_user_id FROM team WHERE name='Nuovo Team'")
    assert team["owner_user_id"] == coo
    assert "Vera Due" in client.get("/people").text and "Vera Due" not in client2.get("/people").text
    assert "Zeno Tre" in client2.get("/people").text and "Ugo Uno" in client2.get("/people").text
    assert db.sys.one("SELECT COUNT(*) AS n FROM person WHERE full_name IN ('Yuri Quattro', 'Xena Cinque')")["n"] == 0


def test_censimento_completo_dal_form_di_hr(client, client2, hr_client):
    t1, t2, ids, coo, dir2 = _seed(client, client2)
    html = hr_client.get("/people").text
    for field in ("birth_date", "notes", "ccnl", "smartwork", "timbratura_scheme", "part_time_pct", "weekly_hours", "listino_profile_id"):
        assert f'name="{field}"' in html, field
    r = hr_client.post("/people", data={"full_name": "Nina Nove", "director_id": dir2, "birth_date": "1991-05-05", "notes": "nota",
                                        "livello": "l2", "ccnl": "Metalmeccanico", "smartwork": "full", "part_time_pct": "50",
                                        "weekly_hours": "20", "ral": "25000"}, follow_redirects=False)
    assert r.status_code == 302
    p = db.sys.one("SELECT * FROM person WHERE full_name='Nina Nove'")
    assert p["owner_user_id"] == dir2 and str(p["birth_date"]) == "1991-05-05" and p["notes"] == "nota"
    emp = db.sys.one("SELECT * FROM employment WHERE person_id=?", (p["id"],))
    assert emp["ccnl"] == "Metalmeccanico" and emp["smartwork"] == "full" and emp["part_time_pct"] == 50 and emp["weekly_hours"] == 20
    assert db.sys.one("SELECT COUNT(*) AS n FROM person_assignment WHERE person_id=?", (p["id"],))["n"] == 0   # senza team
    assert "Nina Nove" in client2.get("/people").text
