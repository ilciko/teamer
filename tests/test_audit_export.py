"""Export dell'audit (0.45): CSV ed Excel con gli stessi filtri della pagina, solo per l'admin."""
import io

from openpyxl import load_workbook

from app import db


def test_export_audit_csv_e_xlsx(client, admin_client, hr_client):
    client.post("/org/team", data={"name": "Alfa", "parent_team_id": "", "lead_person_id": ""})
    hr_client.post("/hr/acting", data={"director_id": "", "next": "/org"})
    html = admin_client.get("/admin?tab=audit").text
    assert 'href="/admin/audit/export.csv?' in html and 'href="/admin/audit/export.xlsx?' in html
    r = admin_client.get("/admin/audit/export.csv?user=coo")
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/csv")
    assert 'filename="teamer-audit-' in r.headers["content-disposition"]
    body = r.text.lstrip("\ufeff")
    lines = body.strip().splitlines()
    assert lines[0] == "quando;utente;evento;dettaglio;ip"
    assert all(";coo;" in l for l in lines[1:]) and any("POST /org/team" in l for l in lines)
    assert not any(";hr1;" in l for l in lines)
    r = admin_client.get("/admin/audit/export.xlsx?event=azione")
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.content))["Audit"]
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == ("Quando", "Utente", "Evento", "Dettaglio", "IP")
    assert len(rows) > 1 and all(row[2] == "azione" for row in rows[1:])
    # tutto il filtro, non la pagina; l'export stesso finisce nell'audit; solo l'admin
    n = db.sys.one("SELECT COUNT(*) AS n FROM audit_event WHERE event='azione'")["n"]
    assert len(rows) - 1 >= n - 1
    assert db.sys.one("SELECT detail FROM audit_event WHERE event='audit_export' ORDER BY id LIMIT 1")["detail"].startswith("csv, ")
    assert client.get("/admin/audit/export.csv").status_code == 403
    assert hr_client.get("/admin/audit/export.csv").status_code == 403
    assert admin_client.get("/admin/audit/export.pdf", follow_redirects=False).status_code == 302
